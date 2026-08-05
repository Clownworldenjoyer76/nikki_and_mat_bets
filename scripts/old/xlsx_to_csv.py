#!/usr/bin/env python3
import sys
import csv
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "docs" / "data"

def die(msg, code=78):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(code)

if len(sys.argv) != 4:
    die("Usage: xlsx_to_csv.py <folder> <week> <season>")

folder = sys.argv[1].strip().lower()
week = int(sys.argv[2])
season = int(sys.argv[3])

if folder not in {"scores","picks","final"}:
    die("Folder must be scores|picks|final")
if not (1 <= week <= 23):
    die("Week must be 1..23")

xlsx = DATA_DIR / folder / f"{season}_wk{week:02d}_{folder}.xlsx"
if not xlsx.exists():
    die(f"Missing XLSX file: {xlsx}")

csvp = xlsx.with_suffix(".csv")
wb = load_workbook(filename=xlsx, read_only=True, data_only=True)
ws = wb.active

rows = []
for row in ws.iter_rows(values_only=True):
    rows.append([("" if v is None else str(v)) for v in row])

with csvp.open("w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerows(rows)

print(csvp)
